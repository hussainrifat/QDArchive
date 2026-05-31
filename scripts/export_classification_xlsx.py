"""
Part 2 Step 4c — Export classification table to XLSX
Columns: repository_id, project_type, project_title,
         primary_class, secondary_class, no_project_files
"""

import sqlite3
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DB_PATH = Path(__file__).parent.parent / "23025313-sq26-classification.db"
OUT_PATH = Path(__file__).parent.parent / "export" / "23025313-sq26-classification.xlsx"


def run():
    conn = sqlite3.connect(str(DB_PATH))

    rows = conn.execute("""
        SELECT
            r.name                          AS repository_id,
            p.type                          AS project_type,
            p.title                         AS project_title,
            p.isic_section_code || ' - ' || p.isic_division_code || ' ' || p.isic_division_name
                                            AS primary_class,
            NULL                            AS secondary_class,
            (SELECT COUNT(*) FROM FILES f WHERE f.project_id = p.id)
                                            AS no_project_files
        FROM PROJECTS p
        JOIN REPOSITORIES r ON r.id = p.repository_id
        ORDER BY r.name, p.type, p.title
    """).fetchall()

    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classification"

    headers = [
        "repository_id", "project_type", "project_title",
        "primary_class", "secondary_class", "no_project_files"
    ]

    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    col_widths = [14, 18, 70, 45, 20, 16]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    ws.freeze_panes = "A2"

    OUT_PATH.parent.mkdir(exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print(f"Rows: {len(rows)} projects")


if __name__ == "__main__":
    run()
